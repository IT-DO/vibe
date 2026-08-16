"""Чтение Excel-выгрузки ЕГРЗ (`PublicRegistrationBook/excelDataFile`).

Реестр умеет отдавать сам себя файлом .xlsx — это самый простой способ забрать
данные: ссылку можно открыть в браузере, без разбора JSON-контракта. Модуль
принимает такой файл (или прямую ссылку на него) и превращает в те же
нормализованные записи, что и API-источник.

Заголовки колонок в выгрузке русские и со временем меняются, поэтому они не
зашиты списком, а сопоставляются по ключевым словам: «Наименование экспертной
организации», «Экспертная организация» и «Организация, проводившая экспертизу»
попадут в одно поле. Колонки, которые сопоставить не удалось, не выбрасываются
молча — они попадают в отчёт :attr:`ExcelSource.report`.
"""

from __future__ import annotations

import re
import time
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterator
from urllib.parse import urlencode

from .models import Conclusion
from .normalize import build_conclusion


def _norm_header(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).lower().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


@dataclass(frozen=True, slots=True)
class HeaderRule:
    """Правило сопоставления заголовка колонки с полем модели."""

    target: str
    require: tuple[str, ...]              # все слова должны встретиться
    exclude: tuple[str, ...] = ()         # ни одного из этих быть не должно

    def matches(self, header: str) -> bool:
        return (
            all(word in header for word in self.require)
            and not any(word in header for word in self.exclude)
        )


#: Порядок значим: узкие правила стоят выше широких, иначе «ИНН экспертной
#: организации» будет поглощено правилом для самой организации.
HEADER_RULES: tuple[HeaderRule, ...] = (
    HeaderRule("id", ("идентификатор",)),

    HeaderRule("organization_inn", ("инн",), ("застройщик", "заказчик", "заявител")),
    HeaderRule("developer_inn", ("инн",)),

    HeaderRule("registry_number", ("номер", "реестр")),
    HeaderRule("registry_number", ("регистрационный", "номер")),
    HeaderRule("conclusion_number", ("номер", "заключен")),
    HeaderRule("conclusion_number", ("номер",), ("объект", "дело")),

    HeaderRule("date_registered", ("дата", "реестр")),
    HeaderRule("date_registered", ("дата", "включен")),
    HeaderRule("date_registered", ("дата", "регистрац")),
    HeaderRule("date_issued", ("дата", "заключен")),
    HeaderRule("date_issued", ("дата", "экспертиз")),
    HeaderRule("date_issued", ("дата",), ("объект",)),

    # «Форма экспертизы» — формулировка из реальной выгрузки ЕГРЗ;
    # «вид»/«тип» держим на случай, если формулировка изменится.
    HeaderRule("expertise_type", ("форма", "экспертиз")),
    HeaderRule("expertise_type", ("вид", "экспертиз"), ("предмет", "объект", "результат")),
    HeaderRule("expertise_type", ("тип", "экспертиз")),
    HeaderRule("result", ("результат",)),
    HeaderRule("result", ("вид", "заключен")),
    HeaderRule("subject_matter", ("предмет",)),
    HeaderRule("subject_matter", ("объект", "экспертиз")),
    HeaderRule("is_repeat", ("повторн",)),

    HeaderRule("object_name", ("наименование", "объект")),
    HeaderRule("object_name", ("объект", "капитальн")),
    HeaderRule("purpose", ("назначение",)),
    # Исключаем «Место нахождения и адрес экспертной организации» — это адрес
    # организации, а не объекта; в реальной выгрузке ЕГРЗ адрес объекта уже
    # встроен текстом в колонку названия объекта, отдельной колонки для него нет.
    HeaderRule("address", ("адрес",), ("организации",)),
    HeaderRule("address", ("место", "располож"), ("организации",)),
    # Код субъекта — раньше общего правила по «субъект», иначе «код субъекта
    # рф» просто не дойдёт до этого правила, если «Субъект РФ» в файле нет.
    HeaderRule("region_code", ("код", "субъект")),
    HeaderRule("region", ("субъект",)),
    HeaderRule("region", ("регион",)),

    HeaderRule("organization", ("экспертн", "организац")),
    HeaderRule("organization", ("организац", "проводивш")),
    HeaderRule("organization", ("организац",), ("застройщик", "заказчик")),
    HeaderRule("developer", ("застройщик",)),
    HeaderRule("developer", ("технический", "заказчик")),
    HeaderRule("developer", ("заявител",)),

    HeaderRule("cost", ("сметн", "стоимост")),
    HeaderRule("cost", ("стоимост",)),
)


def map_headers(headers: list[Any]) -> tuple[dict[int, str], list[str]]:
    """Сопоставляет заголовки с полями модели.

    Возвращает ``(индекс колонки -> поле, список нераспознанных заголовков)``.
    Поле занимает первая подошедшая колонка: в выгрузке встречаются похожие
    заголовки, и дубль не должен затирать уже найденное значение.
    """
    mapping: dict[int, str] = {}
    taken: set[str] = set()
    unmapped: list[str] = []

    for index, raw in enumerate(headers):
        header = _norm_header(raw)
        if not header:
            continue
        for rule in HEADER_RULES:
            if rule.target in taken or not rule.matches(header):
                continue
            mapping[index] = rule.target
            taken.add(rule.target)
            break
        else:
            unmapped.append(str(raw).strip())

    return mapping, unmapped


def _score_header_row(row: list[Any]) -> int:
    """Насколько строка похожа на строку заголовков."""
    mapping, _ = map_headers(row)
    return len(mapping)


#: Верхняя граница колонок для потокового чтения (см. пояснение ниже).
#: Ни одна реальная таблица реестра в этот запас не упрётся.
_SAFE_MAX_COL = 100


def read_rows(
    path: str | Path, *, sheet: str | None = None, probe_rows: int = 12
) -> tuple[list[dict[str, Any]], list[str]]:
    """Читает .xlsx и возвращает ``(записи, нераспознанные заголовки)``.

    Строка заголовков ищется среди первых строк: в выгрузке над таблицей
    попадаются заголовок отчёта и пустые строки.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - зависимость объявлена в pyproject
        raise RuntimeError(
            "Для чтения Excel нужен openpyxl: pip install openpyxl"
        ) from exc

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.worksheets[0]

    try:
        # Выгрузка ЕГРЗ пишет в лист некорректный тег <dimension> (он указывает
        # только на первую колонку, хотя реальных колонок в разы больше — баг
        # генератора файла на стороне сервера). Потоковый iter_rows() в
        # read_only-режиме доверяет этому тегу и без явного max_col обрезает
        # каждую строку до одной ячейки, молча теряя все данные, кроме id.
        # Задаём щедрый max_col сами — лишние пустые колонки после реальных
        # данных потом отфильтруют map_headers() (пустой заголовок) и emit()
        # (пустое значение), так что запас с большим запасом безопасен.
        rows = worksheet.iter_rows(max_col=_SAFE_MAX_COL, values_only=True)
        head: list[list[Any]] = []
        for row in rows:
            head.append(list(row))
            if len(head) >= probe_rows:
                break

        if not head:
            return [], []

        best_index = max(range(len(head)), key=lambda i: _score_header_row(head[i]))
        mapping, unmapped = map_headers(head[best_index])
        if not mapping:
            raise RuntimeError(
                "В файле не найдено ни одной узнаваемой колонки. "
                "Проверьте, что это выгрузка реестра ЕГРЗ, а не другой отчёт."
            )

        records: list[dict[str, Any]] = []

        def emit(row: list[Any]) -> None:
            item = {
                field_name: row[index]
                for index, field_name in mapping.items()
                if index < len(row) and row[index] not in (None, "")
            }
            if item:
                records.append(item)

        for row in head[best_index + 1:]:
            emit(row)
        for row in rows:
            emit(list(row))

        return records, unmapped
    finally:
        workbook.close()


def download(url: str, target: str | Path, *, timeout: float = 180.0) -> Path:
    """Скачивает файл выгрузки по ссылке."""
    import httpx

    destination = Path(target)
    destination.parent.mkdir(parents=True, exist_ok=True)
    headers = {
        "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, */*",
        "User-Agent": "egrz-daily/0.1",
    }
    with httpx.Client(timeout=timeout, follow_redirects=True, headers=headers) as client:
        with client.stream("GET", url) as response:
            response.raise_for_status()
            with destination.open("wb") as handle:
                for chunk in response.iter_bytes():
                    handle.write(chunk)
    return destination


def download_with_retries(
    url: str,
    target: str | Path,
    *,
    timeout: float = 180.0,
    retries: int = 3,
    backoff: float = 3.0,
) -> Path:
    """То же самое, но с повторами при сетевых сбоях.

    Обход всей истории — это десятки-сотни последовательных запросов; один
    транзитный сбой (обрыв связи, таймаут) не должен ронять весь прогон,
    когда до конца ещё далеко.
    """
    last_error: Exception | None = None
    for attempt in range(retries + 1):
        try:
            return download(url, target, timeout=timeout)
        except Exception as exc:  # noqa: BLE001 — ретраим любую сетевую беду
            last_error = exc
            if attempt == retries:
                break
            time.sleep(backoff * (2**attempt))
    raise RuntimeError(f"Не удалось скачать {url}: {last_error}") from last_error


@dataclass
class ExcelSource:
    """Источник поверх Excel-выгрузки реестра.

    ``source`` — путь к скачанному файлу или прямая ссылка на
    ``.../PublicRegistrationBook/excelDataFile?...``.
    """

    source: str
    sheet: str | None = None
    cache_dir: str | Path = "data/downloads"
    report: dict[str, Any] = field(default_factory=dict)

    def _resolve(self) -> Path:
        if str(self.source).startswith(("http://", "https://")):
            target = Path(self.cache_dir) / "egrz-registry.xlsx"
            return download(self.source, target)
        path = Path(self.source)
        if not path.exists():
            raise RuntimeError(f"Файл не найден: {path}")
        return path

    def iter_conclusions(
        self, *, since: str | None = None, limit: int | None = None
    ) -> Iterator[Conclusion]:
        path = self._resolve()
        raw_records, unmapped = read_rows(path, sheet=self.sheet)

        self.report = {
            "file": str(path),
            "rows": len(raw_records),
            "unmapped_columns": unmapped,
        }

        produced = 0
        for item in raw_records:
            record = build_conclusion(dict(item), raw={})
            # В Excel-выгрузке дата включения в реестр может отсутствовать —
            # тогда инкрементальность опирается на дату заключения.
            anchor = record.date_registered or record.date_issued
            if since and anchor and anchor < since:
                continue
            yield record
            produced += 1
            if limit is not None and produced >= limit:
                return


#: Тот же контроллер, что и в ссылке-примере с сайта — без него не собрать URL.
DEFAULT_BACKFILL_BASE_URL = "https://open-api.egrz.ru/api/PublicRegistrationBook/excelDataFile"


@dataclass
class ExcelBackfillSource:
    """Постраничный обход ВСЕЙ истории реестра (сотни тысяч записей).

    У сервера нет потоковой выдачи — каждая страница это отдельный HTTP-запрос,
    генерирующий отдельный .xlsx-файл на его стороне. Полная история по всей
    стране (700 000+ записей) — это не один файл, а десятки-сотни
    последовательных запросов по ``page_size`` записей, отсортированных по
    дате заключения (``$orderby``) и сдвинутых ``$skip``.

    Спроектировано так, чтобы длинный обход переживал обрыв: временный файл
    каждой страницы удаляется сразу после разбора (не копим сотни МБ на
    диске), записи отдаются по одной генератором (не копим 700 000+ записей
    в памяти), а ``on_page`` вызывается после каждой успешно обработанной
    страницы — CLI использует это, чтобы сохранить позицию в БД и продолжить
    прерванный обход с того же места, а не с нуля.
    """

    base_url: str = DEFAULT_BACKFILL_BASE_URL
    order_by: str = "ExpertiseDate desc"
    page_size: int = 5000
    start_skip: int = 0
    max_pages: int | None = None
    rate_limit: float = 1.0
    cache_dir: str | Path = "data/downloads"
    #: Вызывается как on_page(skip_этой_страницы, число_записей_на_ней).
    on_page: Callable[[int, int], None] | None = None
    report: dict[str, Any] = field(default_factory=dict)

    def _page_url(self, skip: int) -> str:
        # safe="$" — иначе urlencode экранирует его в %24. Сервер, скорее
        # всего, понял бы и так (%24 — стандартная запись для $ в query string),
        # но литеральный $ читаем в логах и совпадает с проверенным вручную
        # примером ссылки, так что незачем полагаться на декодирование.
        query = urlencode(
            {"$orderby": self.order_by, "$count": "true", "$top": self.page_size, "$skip": skip},
            safe="$",
        )
        return f"{self.base_url}?{query}"

    def iter_conclusions(self) -> Iterator[Conclusion]:
        skip = self.start_skip
        pages_done = 0
        total_records = 0
        cache_dir = Path(self.cache_dir)
        cache_dir.mkdir(parents=True, exist_ok=True)

        while self.max_pages is None or pages_done < self.max_pages:
            target = cache_dir / f"backfill-{skip}.xlsx"
            try:
                download_with_retries(self._page_url(skip), target)
                raw_records, unmapped = read_rows(target)
            finally:
                target.unlink(missing_ok=True)

            for item in raw_records:
                yield build_conclusion(dict(item), raw={})

            total_records += len(raw_records)
            pages_done += 1
            self.report = {
                "pages": pages_done,
                "records": total_records,
                "last_skip": skip,
                "unmapped_columns": unmapped,
            }

            if self.on_page:
                self.on_page(skip, len(raw_records))

            # Короче полного размера страницы — это последняя страница истории.
            if len(raw_records) < self.page_size:
                self.report["finished"] = True
                return

            skip += self.page_size
            if self.rate_limit:
                time.sleep(self.rate_limit)

        # Обход остановлен по max_pages — истории могло остаться ещё много.
        self.report["finished"] = False
